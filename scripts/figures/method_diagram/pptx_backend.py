"""PowerPoint back end: every scene element becomes a native, named shape.

Boxes, the brace and the speech bubble are preset shapes (with their yellow
adjustment handles), arrows are connectors, icons and curves are freeforms
(editable with Edit Points), and related shapes are grouped.
"""

from __future__ import annotations

import logging
from collections.abc import Callable
from pathlib import Path

from geometry import callout_adjustments
from icons import load_icon
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from primitives import (
    BLACK,
    WHITE,
    Box,
    Brace,
    Callout,
    Connector,
    Element,
    Ellipse,
    Group,
    Icon,
    PathShape,
    Scene,
    Segment,
    Stroke,
    Text,
    TextLine,
)

logger = logging.getLogger(__name__)

EMU_PER_PT = 12700
PATH_SCALE = 100  # custom-geometry coordinates must be integers: keep 2 decimals
BLANK_LAYOUT = 6
PATH_TAGS = {"M": "a:moveTo", "L": "a:lnTo", "C": "a:cubicBezTo", "Z": "a:close"}


def _emu(value: float) -> Emu:
    return Emu(round(value * EMU_PER_PT))


def _strip_style(shape) -> None:
    """Drop the theme style reference so only explicit formatting applies."""
    for style in shape._element.findall(qn("p:style")):
        shape._element.remove(style)


def _set_stroke(shape, stroke: Stroke | None) -> None:
    if stroke is None:
        shape.line.fill.background()
        return
    shape.line.color.rgb = RGBColor.from_string(stroke.color)
    shape.line.width = Pt(stroke.width)
    if stroke.dashed:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def _set_fill(shape, fill: str | None) -> None:
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)


def _add_arrowhead(shape) -> None:
    tail = etree.SubElement(shape.line._get_or_add_ln(), qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "lg")
    tail.set("len", "lg")


def _set_text(text_frame, lines: tuple[TextLine, ...], color: str,
              inset: tuple[float, float, float, float] = (0.0, 0.0, 0.0, 0.0)) -> None:
    text_frame.word_wrap = False
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    (text_frame.margin_left, text_frame.margin_top,
     text_frame.margin_right, text_frame.margin_bottom) = (_emu(v) for v in inset)
    for i, line in enumerate(lines):
        para = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        for spec in line:
            run = para.add_run()
            run.text = spec.text
            run.font.size = Pt(spec.size)
            run.font.bold = spec.bold
            run.font.italic = spec.italic
            run.font.name = spec.font
            run.font.color.rgb = RGBColor.from_string(color)
            if spec.subscript:
                run._r.get_or_add_rPr().set("baseline", "-25000")
            if spec.spacing:
                run._r.get_or_add_rPr().set("spc", str(round(spec.spacing * 100)))


def _freeform(shapes, segments: tuple[Segment, ...], box: tuple[float, float, float, float],
              path_size: tuple[float, float], open_path: bool):
    """Freeform occupying ``box`` (pt) whose path, of extent ``path_size``, is ``segments``."""
    x, y, w, h = box
    builder = shapes.build_freeform(_emu(x), _emu(y))
    builder.add_line_segments(
        [(_emu(x + w), _emu(y)), (_emu(x + w), _emu(y + h)), (_emu(x), _emu(y + h))], close=True
    )
    shape = builder.convert_to_shape()
    _strip_style(shape)
    path = shape._element.find(".//" + qn("a:path"))
    for child in list(path):
        path.remove(child)
    path.set("w", str(max(1, round(path_size[0] * PATH_SCALE))))
    path.set("h", str(max(1, round(path_size[1] * PATH_SCALE))))
    if open_path:
        path.set("fill", "none")
    for seg in segments:
        op = etree.SubElement(path, qn(PATH_TAGS[seg[0]]))
        for k in range(1, len(seg), 2):
            pt = etree.SubElement(op, qn("a:pt"))
            pt.set("x", str(round(seg[k] * PATH_SCALE)))
            pt.set("y", str(round(seg[k + 1] * PATH_SCALE)))
    return shape


def _box(shapes, el: Box):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if el.radius > 0 else MSO_SHAPE.RECTANGLE
    shape = shapes.add_shape(kind, _emu(el.x), _emu(el.y), _emu(el.w), _emu(el.h))
    _strip_style(shape)
    if el.radius > 0:
        shape.adjustments[0] = el.radius / min(el.w, el.h)
    _set_fill(shape, el.fill)
    _set_stroke(shape, el.stroke)
    _set_text(shape.text_frame, el.lines, BLACK, el.inset)
    return shape


def _text(shapes, el: Text):
    shape = shapes.add_textbox(_emu(el.x), _emu(el.y), _emu(el.w), _emu(el.h))
    _set_text(shape.text_frame, el.lines, el.color)
    return shape


def _ellipse(shapes, el: Ellipse):
    shape = shapes.add_shape(MSO_SHAPE.OVAL, _emu(el.x), _emu(el.y), _emu(el.w), _emu(el.h))
    _strip_style(shape)
    _set_fill(shape, el.fill)
    _set_stroke(shape, None)
    return shape


def _icon_copy(shapes, el: Icon, fill: str, stroke: Stroke | None):
    icon = load_icon(el.key)
    box = (el.x, el.y, el.h * icon.width / icon.height, el.h)
    shape = _freeform(shapes, icon.segments, box, (icon.width, icon.height), open_path=False)
    _set_fill(shape, fill)
    _set_stroke(shape, stroke)
    return shape


def _icon(shapes, el: Icon):
    if el.halo <= 0:
        return _icon_copy(shapes, el, el.fill, None)
    group = shapes.add_group_shape()  # white knockout behind a solid copy
    _icon_copy(group.shapes, el, WHITE, Stroke(WHITE, el.halo)).name = f"{el.name} halo"
    _icon_copy(group.shapes, el, el.fill, None).name = f"{el.name} fill"
    group.shapes._recalculate_extents()
    return group


def _connector(shapes, el: Connector):
    kind = MSO_CONNECTOR.CURVE if el.kind == "curved" else MSO_CONNECTOR.STRAIGHT
    shape = shapes.add_connector(kind, _emu(el.x1), _emu(el.y1), _emu(el.x2), _emu(el.y2))
    _strip_style(shape)
    _set_stroke(shape, el.stroke)
    if el.arrow:
        _add_arrowhead(shape)
    return shape


def _path(shapes, el: PathShape):
    xs = [seg[k] for seg in el.segments for k in range(1, len(seg), 2)]
    ys = [seg[k] for seg in el.segments for k in range(2, len(seg), 2)]
    x0, y0 = min(xs), min(ys)
    w, h = max(xs) - x0, max(ys) - y0
    local = tuple(
        (seg[0], *[v - (x0 if k % 2 else y0) for k, v in enumerate(seg[1:], start=1)])
        for seg in el.segments
    )
    shape = _freeform(shapes, local, (x0, y0, w, h), (w, h), open_path=el.fill is None)
    _set_fill(shape, el.fill)
    _set_stroke(shape, el.stroke)
    if el.arrow:
        _add_arrowhead(shape)
    return shape


def _brace(shapes, el: Brace):
    shape = shapes.add_shape(MSO_SHAPE.RIGHT_BRACE, _emu(el.x), _emu(el.y), _emu(el.w), _emu(el.h))
    _strip_style(shape)
    shape.adjustments[0] = el.radius / min(el.w, el.h)
    shape.adjustments[1] = 0.5
    _set_fill(shape, None)
    _set_stroke(shape, el.stroke)
    return shape


def _callout(shapes, el: Callout):
    shape = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
                             _emu(el.x), _emu(el.y), _emu(el.w), _emu(el.h))
    _strip_style(shape)
    for i, value in enumerate(callout_adjustments(el)):
        shape.adjustments[i] = value
    _set_fill(shape, el.fill)
    _set_stroke(shape, el.stroke)
    _set_text(shape.text_frame, el.lines, BLACK)
    return shape


def _group(shapes, el: Group):
    group = shapes.add_group_shape()
    for child in el.children:
        _add(group.shapes, child)
    group.shapes._recalculate_extents()
    return group


_BUILDERS: dict[type, Callable] = {
    Box: _box, Text: _text, Ellipse: _ellipse, Icon: _icon, Connector: _connector,
    PathShape: _path, Brace: _brace, Callout: _callout, Group: _group,
}


def _add(shapes, el: Element):
    shape = _BUILDERS[type(el)](shapes, el)
    shape.name = el.name
    return shape


def write_pptx(scene: Scene, path: Path) -> None:
    """Write the scene on a single slide sized exactly like the figure."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = _emu(scene.width), _emu(scene.height)
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    for el in scene.elements:
        _add(slide.shapes, el)
    if scene.notes:
        slide.notes_slide.notes_text_frame.text = scene.notes
    prs.save(str(path))
    logger.info("Wrote %s", path)
